"""
CH09 LangChain 연결 전략 — 문서 검색 도구.

@tool 데코레이터를 사용하여 LangChain Agent에서 호출 가능한 도구를 정의합니다.
ChromaDB 벡터 검색을 시도하며, 불가 시 data/docs/ 원본 문서를 파싱하여 키워드 검색합니다.
"""

import os
import logging
from pathlib import Path
from typing import Union

from langchain_core.tools import tool

logger = logging.getLogger(__name__)


def _search_from_chroma(query: str, top_k: int = 3) -> Union[list[dict], None]:
    """ChromaDB에서 의미론적 유사도 검색을 수행합니다.

    Args:
        query: 검색할 질문 또는 키워드
        top_k: 반환할 최대 문서 수

    Returns:
        검색된 문서 딕셔너리 목록 또는 None (ChromaDB 연결 실패 시)
    """
    try:
        import chromadb
        from sentence_transformers import SentenceTransformer

        chroma_dir = os.getenv("CHROMA_PERSIST_DIR", "./data/chroma_db")
        embedding_model = os.getenv("EMBEDDING_MODEL", "jhgan/ko-sroberta-multitask")

        # ① ChromaDB 클라이언트 초기화
        client = chromadb.PersistentClient(path=chroma_dir)
        collection_name = os.getenv("CHROMA_COLLECTION_NAME", "metacoding_documents")
        collection = client.get_collection(collection_name)

        # ② 임베딩 생성
        model = SentenceTransformer(embedding_model)
        query_embedding = model.encode(query).tolist()

        # ③ 유사도 검색
        results = collection.query(
            query_embeddings=[query_embedding],
            n_results=top_k,
            include=["documents", "metadatas", "distances"],
        )

        # ④ 결과 포맷팅
        formatted = []
        if results["documents"] and results["documents"][0]:
            for doc, meta, dist in zip(
                results["documents"][0],
                results["metadatas"][0],
                results["distances"][0],
            ):
                formatted.append({
                    "content": doc,
                    "source": meta.get("source", "unknown"),
                    "score": round(1.0 - dist, 4),
                })
        return formatted

    except Exception as exc:
        logger.warning("ChromaDB 검색 실패, data/docs/ 키워드 검색 사용: %s", exc)
        return None


def _parse_docs_dir(docs_dir: Path) -> list[dict]:
    """data/docs/의 원본 문서를 파싱하여 검색 가능한 딕셔너리 목록으로 반환합니다.

    Args:
        docs_dir: 원본 문서가 위치한 디렉토리 경로

    Returns:
        {"content": str, "source": str} 형태의 딕셔너리 리스트
    """
    import pypdf
    from docx import Document as DocxDocument
    import openpyxl

    docs: list[dict] = []
    if not docs_dir.exists():
        return docs

    for file_path in sorted(docs_dir.rglob("*")):
        suffix = file_path.suffix.lower()
        source = file_path.stem

        if suffix == ".pdf":
            try:
                with open(file_path, "rb") as f:
                    reader = pypdf.PdfReader(f)
                    for page in reader.pages:
                        text = (page.extract_text() or "").strip()
                        if text and len(text) > 30:
                            docs.append({"content": text, "source": source})
            except Exception:
                pass
        elif suffix == ".docx":
            try:
                doc = DocxDocument(str(file_path))
                text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
                if text:
                    docs.append({"content": text, "source": source})
            except Exception:
                pass
        elif suffix == ".xlsx":
            try:
                wb = openpyxl.load_workbook(str(file_path), data_only=True)
                for name in wb.sheetnames:
                    ws = wb[name]
                    rows = []
                    for row in ws.iter_rows():
                        cells = [str(c.value).strip() for c in row if c.value is not None]
                        if cells:
                            rows.append(" | ".join(cells))
                    if rows:
                        docs.append({"content": "\n".join(rows), "source": source})
            except Exception:
                pass

    return docs


def _search_from_mock(query: str, top_k: int = 3) -> list[dict]:
    """data/docs/ 원본 문서를 파싱하여 키워드 기반 검색을 수행합니다.

    Args:
        query: 검색할 질문 또는 키워드
        top_k: 반환할 최대 문서 수

    Returns:
        검색된 문서 딕셔너리 목록 (score 내림차순 정렬)
    """
    docs_dir = Path(__file__).resolve().parent.parent.parent / "data" / "docs"
    parsed_docs = _parse_docs_dir(docs_dir)

    if not parsed_docs:
        return []

    query_lower = query.lower()
    scored_docs = []
    for doc in parsed_docs:
        score = 0.0
        for word in query_lower.split():
            if word in doc["content"].lower():
                score += 0.1
        scored_docs.append({**doc, "score": round(min(score + 0.5, 1.0), 4)})

    scored_docs.sort(key=lambda x: x["score"], reverse=True)
    return scored_docs[:top_k]


# --- INPUT ---
@tool
def search_documents(query: str) -> list[dict]:
    """사내 규정, 가이드라인, 정책 등 비정형 문서 내용을 검색합니다.

    휴가 규정, 보안 수칙, 온보딩 가이드, 비용 정산 기준 등
    사내 문서에서 관련 내용을 의미론적 유사도 기반으로 찾아 반환합니다.

    Args:
        query: 검색할 질문 또는 키워드 (예: "연차 사용 규정", "재택근무 조건")

    Returns:
        관련 문서 내용, 출처 파일명, 유사도 점수가 담긴 딕셔너리 목록.
        점수가 높을수록 질문과 관련도가 높습니다.
    """
    # --- PROCESS ---
    logger.info("[search_documents] 검색 쿼리: %s", query)

    # ChromaDB 검색 시도 → 실패 시 data/docs/ 키워드 검색으로 폴백
    results = _search_from_chroma(query)
    if results is None:
        results = _search_from_mock(query)

    # --- OUTPUT ---
    logger.info("[search_documents] 검색 결과 수: %d", len(results))
    return results
